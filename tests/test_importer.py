"""Importer: diverse workbook/CSV layouts must map to the right table with values unchanged."""
import io
import sys
from datetime import datetime
from pathlib import Path

import numpy as np
import pandas as pd
import pytest
from openpyxl import Workbook

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))

from core.importer import import_file  # noqa: E402

DATES = [datetime(2024, m, 28) for m in range(1, 7)]
FUND = [0.012, -0.0045, 0.0301, 0.0, -0.0123, 0.0077]
MKT = [0.01, -0.002, 0.025, 0.004, -0.015, 0.006]


def xlsx(build) -> bytes:
    wb = Workbook()
    wb.remove(wb.active)
    build(wb)
    b = io.BytesIO()
    wb.save(b)
    return b.getvalue()


def check_returns(r):
    assert list(r.df.columns[:3]) == ["Date", "Fund", "Market"]
    assert r.df["Fund"].tolist() == FUND           # exact, not approximately
    assert r.df["Market"].tolist() == MKT
    assert pd.to_datetime(r.df["Date"]).tolist() == [pd.Timestamp(d) for d in DATES]


def test_clean_table():
    def b(wb):
        ws = wb.create_sheet("Data")
        ws.append(["Date", "Fund", "Market"])
        for row in zip(DATES, FUND, MKT):
            ws.append(list(row))
    r = import_file(xlsx(b), "a.xlsx")
    check_returns(r)
    assert r.header_row == 1 and r.first_row == 2 and r.last_row == 7 and not r.needs_confirmation


def test_title_rows_notes_blank_columns_and_merged_heading():
    def b(wb):
        ws = wb.create_sheet("Report")
        ws["B2"] = "Quarterly Fund Performance"
        ws["B3"] = "Prepared by Research"
        ws["C5"] = "Returns"
        ws.merge_cells("C5:D5")          # merged group heading above the header
        ws["B6"], ws["C6"], ws["D6"] = "Date", "Fund", "Market"
        for i, row in enumerate(zip(DATES, FUND, MKT)):
            for j, v in enumerate(row):
                ws.cell(row=7 + i, column=2 + j, value=v)
        ws["B14"] = "Source: internal records"
        ws["B15"] = "Returns are net of fees"
    r = import_file(xlsx(b), "b.xlsx")
    check_returns(r)
    assert r.header_row == 6 and r.first_row == 7 and r.last_row == 12
    assert r.first_col == 2
    assert any("title" in n for n in r.notes) and any("note" in n or "footer" in n for n in r.notes)


def test_multi_sheet_asks_for_confirmation_and_cover_sheet_ignored():
    def b(wb):
        ws = wb.create_sheet("Cover")
        ws["A1"] = "Confidential model"
        ws["A3"] = "See Data tab"
        ws = wb.create_sheet("Data")
        ws.append(["Date", "Fund", "Market"])
        for row in zip(DATES, FUND, MKT):
            ws.append(list(row))
        ws = wb.create_sheet("Old")
        ws.append(["Year", "A", "B"])
        for y in range(2015, 2019):
            ws.append([y, y * 0.1, y * 0.2])
    data = xlsx(b)
    r = import_file(data, "c.xlsx")
    assert r.sheet == "Data" and r.needs_confirmation and "Old" in r.question
    check_returns(r)
    r2 = import_file(data, "c.xlsx", sheet="Old")
    assert r2.sheet == "Old" and not r2.needs_confirmation and r2.df["A"].tolist()[0] == 2015 * 0.1


def test_sideways_table_with_periods_across_columns():
    periods = ["FY2019", "FY2020", "FY2021", "FY2022", "FY2023"]
    rev = [100.0, 112.5, 98.0, 130.25, 141.0]
    margin = [0.21, 0.19, 0.15, 0.22, 0.24]

    def b(wb):
        ws = wb.create_sheet("P&L")
        ws.append(["Company X summary"])
        ws.append([])
        ws.append(["Metric"] + periods)
        ws.append(["Revenue"] + rev)
        ws.append(["EBITDA margin"] + margin)
    r = import_file(xlsx(b), "d.xlsx")
    assert r.transposed
    assert list(r.df.columns) == ["Period", "Revenue", "EBITDA margin"]
    assert r.df["Revenue"].tolist() == rev and r.df["EBITDA margin"].tolist() == margin
    assert r.df["Period"].tolist()[0] == pd.Timestamp(2019, 12, 31)


def test_sideways_quarters_and_override_off():
    def b(wb):
        ws = wb.create_sheet("Q")
        ws.append(["", "Q1 2023", "Q2 2023", "Q3 2023", "Q4 2023"])
        ws.append(["NOI", 10, 11, 12, 13])
        ws.append(["Occupancy", 0.91, 0.92, 0.9, 0.95])
    data = xlsx(b)
    r = import_file(data, "e.xlsx")
    assert r.transposed and r.df["NOI"].tolist() == [10, 11, 12, 13]
    assert r.df["Period"].tolist()[-1] == pd.Timestamp(2023, 12, 31)
    r2 = import_file(data, "e.xlsx", sheet="Q", transpose=False)
    assert not r2.transposed and "Q1 2023" in r2.df.columns


def test_csv_with_title_lines_and_semicolons():
    raw = "Loan tape export\nGenerated 2026-10-01\n\nloan_id;spread_bps;leverage\nA1;450;4.5\nA2;525;5.25\nA3;610;6\n"
    r = import_file(raw.encode(), "f.csv")
    assert list(r.df.columns) == ["loan_id", "spread_bps", "leverage"]
    assert r.df["spread_bps"].tolist() == [450, 525, 610] and r.df["leverage"].tolist() == [4.5, 5.25, 6.0]
    assert r.header_row == 4


def test_csv_text_formatted_numbers_kept_exactly():
    raw = 'Date,Fund,Assets\n2024-01-31,1.5%,"1,200"\n2024-02-29,(0.5%),"1,350"\n2024-03-31,2%,"1,400"\n'
    r = import_file(raw.encode(), "g.csv")
    assert r.df["Fund"].tolist() == ["1.5%", "(0.5%)", "2%"]   # left as typed; converted later, visibly
    assert r.df["Assets"].tolist() == ["1,200", "1,350", "1,400"]


def test_no_header_row_requires_confirmation():
    raw = "1,2,3\n4,5,6\n7,8,9\n10,11,12\n"
    r = import_file(raw.encode(), "h.csv")
    assert r.needs_confirmation and list(r.df.columns) == ["Column 1", "Column 2", "Column 3"]
    assert r.df["Column 1"].tolist() == [1, 4, 7, 10]


def test_header_override():
    def b(wb):
        ws = wb.create_sheet("S")
        ws.append(["Units: USD m", "", ""])
        ws.append(["Year", "Sales", "Cost"])
        for y in range(2018, 2024):
            ws.append([y, y - 1900.0, (y - 1900) / 2])
    data = xlsx(b)
    r = import_file(data, "i.xlsx")
    assert list(r.df.columns) == ["Year", "Sales", "Cost"]
    r2 = import_file(data, "i.xlsx", sheet="S", header_row=2, last_row=5)
    assert r2.df["Sales"].tolist() == [118.0, 119.0, 120.0]


def test_unusable_file_explains():
    with pytest.raises(ValueError, match="could not find a table"):
        import_file(b"just some text\nmore text\n", "j.csv")


def test_demo_b_round_trip_values_exact():
    from core.demo_data import make_private_credit
    df = make_private_credit()
    b = io.BytesIO()
    df.to_excel(b, index=False, sheet_name="Loans")
    r = import_file(b.getvalue(), "loans.xlsx")
    for c in ["spread_bps", "leverage_x", "ebitda_musd"]:
        np.testing.assert_array_equal(r.df[c].to_numpy(dtype=float), df[c].to_numpy(dtype=float))
    assert r.df["sector"].tolist() == df["sector"].tolist()


def test_sideways_integer_years():
    def b(wb):
        ws = wb.create_sheet("Y")
        ws.append(["Metric", 2020, 2021, 2022, 2023])
        ws.append(["Revenue", 50, 55, 61, 70])
        ws.append(["Capex", 5, 7, 6, 9])
    r = import_file(xlsx(b), "k.xlsx")
    assert r.transposed and r.df["Capex"].tolist() == [5, 7, 6, 9]
    assert r.df["Period"].tolist()[0] == pd.Timestamp(2020, 12, 31)
